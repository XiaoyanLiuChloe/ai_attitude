# Import necessary libraries
import numpy as np
import pandas as pd
import seaborn as sns
import matplotlib.pyplot as plt
from sklearn.preprocessing import StandardScaler
from sklearn.decomposition import PCA
from sklearn.ensemble import RandomForestRegressor
from sklearn.model_selection import GridSearchCV, KFold, cross_val_score
from sklearn.metrics import r2_score, mean_squared_error, mean_absolute_error
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import io
from tabulate import tabulate
from colorama import Fore, Back, Style, init

# Initialize colorama
init()

# Create output directories
os.makedirs('outputs', exist_ok=True)
os.makedirs('outputs/tables', exist_ok=True)
os.makedirs('outputs/visualizations', exist_ok=True)

# Create a new Word document
doc = Document()
doc.add_heading('AI Attitude Analysis Report', 0)

# Set style for visualizations
plt.style.use('seaborn-v0_8-whitegrid')
colors = sns.color_palette("viridis", 10)

# Function to create safe filename
def create_safe_filename(text):
    safe_name = "".join(c if c.isalnum() or c == '_' else '_' for c in text)
    while '__' in safe_name:
        safe_name = safe_name.replace('__', '_')
    return safe_name.strip('_')

# Function to add a section title to the document
def add_section_title(title):
    doc.add_heading(title, level=1)
    print("\n" + "="*80)
    print(" "*30 + Fore.CYAN + Style.BRIGHT + title + Style.RESET_ALL)
    print("="*80 + "\n")

# Function to add a table to the document
def add_table_to_doc(data, headers, title):
    print(Fore.GREEN + Style.BRIGHT + f"\n{title}" + Style.RESET_ALL)
    print(tabulate(data, headers=headers, tablefmt="fancy_grid", floatfmt=".4f"))
    
    doc.add_heading(title, level=2)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = str(header)
        header_cells[i].paragraphs[0].runs[0].bold = True
    
    for row_data in data:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = f"{cell_data:.4f}" if isinstance(cell_data, float) else str(cell_data)
    
    doc.add_paragraph()

# Function to add a figure to the document
def add_figure_to_doc(fig, title, filename):
    fig.savefig(f'outputs/visualizations/{filename}.png', dpi=300, bbox_inches='tight')
    doc.add_heading(title, level=2)
    doc.add_picture(f'outputs/visualizations/{filename}.png', width=Inches(6))
    doc.add_paragraph()
    plt.close()

# Load the dataset
add_section_title("DATA LOADING AND PREPROCESSING")
try:
    data = pd.read_csv('/Users/yuyan/Library/CloudStorage/OneDrive-CUHK-Shenzhen/CSS5230ML/project/survey_data.csv')
    print(Fore.GREEN + f"Successfully loaded dataset with {data.shape[0]} rows and {data.shape[1]} columns" + Style.RESET_ALL)
    doc.add_paragraph(f"Successfully loaded dataset with {data.shape[0]} rows and {data.shape[1]} columns")
except Exception as e:
    print(Fore.RED + f"Error loading dataset: {str(e)}" + Style.RESET_ALL)
    raise

# Define variables for comfortableness and capability
com_vars = [f'Scom{i}' for i in range(1, 43)]
cap_vars = [f'Scap{i}' for i in range(1, 43)]

# Filter to only include available variables
com_vars = [var for var in com_vars if var in data.columns]
cap_vars = [var for var in cap_vars if var in data.columns]

print(Fore.YELLOW + f"Found {len(com_vars)} comfortableness variables and {len(cap_vars)} capability variables" + Style.RESET_ALL)

# Function to perform PCA analysis
def perform_pca_analysis(data, variables, title_prefix):
    add_section_title(f"PCA ANALYSIS: {title_prefix.upper()}")
    
    # Extract and standardize data
    X = data[variables].dropna()
    print(Fore.YELLOW + f"Analyzing {len(variables)} variables with {X.shape[0]} valid samples" + Style.RESET_ALL)
    
    scaler = StandardScaler()
    X_scaled = scaler.fit_transform(X)
    
    # Perform PCA
    pca = PCA(n_components=2)
    X_pca = pca.fit_transform(X_scaled)
    
    # Create scree plot
    plt.figure(figsize=(10, 6))
    components = np.arange(1, 3)
    plt.bar(components, pca.explained_variance_ratio_, alpha=0.7, color=colors)
    plt.plot(components, np.cumsum(pca.explained_variance_ratio_), 'o-', color='red', linewidth=2)
    plt.title(f'{title_prefix} Scree Plot')
    plt.xlabel('Principal Component')
    plt.ylabel('Explained Variance Ratio')
    add_figure_to_doc(plt.gcf(), f'{title_prefix} Scree Plot', f'{title_prefix.lower()}_scree_plot')
    
    # Create loadings plot
    loadings = pd.DataFrame(pca.components_.T, 
                          columns=['Component 1', 'Component 2'], 
                          index=variables)
    
    plt.figure(figsize=(12, 8))
    plt.scatter(loadings['Component 1'], loadings['Component 2'], s=100, color=colors[0], alpha=0.7)
    for i, (var, loading) in enumerate(loadings.iterrows()):
        plt.arrow(0, 0, loading['Component 1'], loading['Component 2'], 
                  head_width=0.02, head_length=0.02, fc='blue', ec='blue', alpha=0.5)
        plt.text(loading['Component 1']*1.1, loading['Component 2']*1.1, var, 
                 color='darkblue', ha='center', va='center', fontsize=9)
    
    plt.title(f'{title_prefix} Loadings Plot')
    add_figure_to_doc(plt.gcf(), f'{title_prefix} Loadings Plot', f'{title_prefix.lower()}_loadings_plot')
    
    return X_pca, pca, loadings

# Perform PCA for both comfortableness and capability
try:
    com_pca, com_pca_model, com_loadings = perform_pca_analysis(data, com_vars, "Comfortableness")
    cap_pca, cap_pca_model, cap_loadings = perform_pca_analysis(data, cap_vars, "Capability")
except Exception as e:
    print(Fore.RED + f"Error in PCA analysis: {str(e)}" + Style.RESET_ALL)
    raise

# Create final DataFrame with PCA factors
pca_factors = pd.DataFrame({
    'ScomFactor1': com_pca[:, 0],
    'ScomFactor2': com_pca[:, 1],
    'ScapFactor1': cap_pca[:, 0],
    'ScapFactor2': cap_pca[:, 1]
}, index=data.index)

# Save PCA factors
pca_factors.to_csv('outputs/pca_factors.csv')
print(Fore.GREEN + "PCA factors saved to outputs/pca_factors.csv" + Style.RESET_ALL)

# Function to perform random forest regression
def perform_rf_analysis(X, y, title_prefix):
    add_section_title(f"RANDOM FOREST ANALYSIS: {title_prefix.upper()}")
    
    # Define hyperparameter grid
    param_grid = {
        'n_estimators': [100, 200],
        'max_depth': [None, 10],
        'min_samples_split': [2, 5],
        'max_features': ['sqrt', 'log2']
    }
    
    # Create KFold object
    kf = KFold(n_splits=5, shuffle=True, random_state=42)
    
    # Initialize GridSearchCV
    grid_search = GridSearchCV(
        estimator=RandomForestRegressor(random_state=42),
        param_grid=param_grid,
        cv=kf,
        scoring='neg_mean_squared_error',
        n_jobs=-1
    )
    
    # Fit model
    grid_search.fit(X, y)
    best_rf = grid_search.best_estimator_
    
    # Calculate performance metrics
    y_pred_train = best_rf.predict(X)
    r2_train = r2_score(y, y_pred_train)
    mse_train = mean_squared_error(y, y_pred_train)
    mae_train = mean_absolute_error(y, y_pred_train)
    
    # Cross-validation performance
    cv_scores = cross_val_score(best_rf, X, y, cv=kf, scoring='neg_mean_squared_error')
    mse_cv = -np.mean(cv_scores)
    
    # Get feature importance
    importances = best_rf.feature_importances_
    
    # Plot feature importance
    plt.figure(figsize=(10, 6))
    plt.bar(X.columns, importances)
    plt.title(f'Feature Importance for {title_prefix}')
    plt.xticks(rotation=45)
    add_figure_to_doc(plt.gcf(), f'{title_prefix} Feature Importance', 
                     f'{title_prefix.lower()}_feature_importance')
    
    return {
        'best_params': grid_search.best_params_,
        'r2_train': r2_train,
        'mse_train': mse_train,
        'mae_train': mae_train,
        'mse_cv': mse_cv,
        'importances': importances
    }

# Perform regression analysis for each attitude item
add_section_title("REGRESSION ANALYSIS")
results = []

# Define dependent variables
pr_cols = [f'PR{i}' for i in range(1, 17)]
nr_cols = [f'NR{i}' for i in range(1, 17)]
all_cols = pr_cols + nr_cols

for col in all_cols:
    if col not in data.columns:
        print(Fore.YELLOW + f"Warning: {col} not found in dataset" + Style.RESET_ALL)
        continue
    
    print(Fore.CYAN + f"\nProcessing {col}" + Style.RESET_ALL)
    try:
        y = data[col]
        rf_results = perform_rf_analysis(pca_factors, y, col)
        
        results.append([
            col,
            rf_results['best_params'],
            rf_results['r2_train'],
            rf_results['mse_train'],
            rf_results['mae_train'],
            rf_results['mse_cv'],
            *rf_results['importances']
        ])
    except Exception as e:
        print(Fore.RED + f"Error processing {col}: {str(e)}" + Style.RESET_ALL)
        continue

# Create results DataFrame
results_df = pd.DataFrame(results, columns=[
    'DV', 'Best_Params',
    'R2_train', 'MSE_train', 'MAE_train', 'MSE_cv',
    'ScomFactor1_importance', 'ScomFactor2_importance',
    'ScapFactor1_importance', 'ScapFactor2_importance'
])

# Save results
results_df.to_csv('outputs/rf_results.csv', index=False)
print(Fore.GREEN + "Regression results saved to outputs/rf_results.csv" + Style.RESET_ALL)

# Add results to document
add_table_to_doc(results_df.values, results_df.columns, "Random Forest Regression Results")

# Save the document
doc.save('outputs/AI_Attitude_Analysis_Report.docx')
print(Fore.GREEN + Style.BRIGHT + "\nAnalysis completed successfully! Report saved as 'AI_Attitude_Analysis_Report.docx'" + Style.RESET_ALL) 