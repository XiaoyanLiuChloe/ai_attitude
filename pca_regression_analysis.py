import pandas as pd
import numpy as np
from sklearn.ensemble import RandomForestRegressor
from sklearn.metrics import r2_score, mean_squared_error, mean_absolute_error
from sklearn.model_selection import GridSearchCV, KFold, cross_val_score
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

# Load PCA results and original data
pca_factors = pd.read_csv('data/pca_factors.csv')
data = pd.read_csv('data/survey_data.csv')

# Prepare predictors (4 principal components)
X = pca_factors[['ScomFactor1', 'ScomFactor2', 'ScapFactor1', 'ScapFactor2']]

# Define dependent variables (PR1-16 and NR1-16)
pr_cols = [f'PR{i}' for i in range(1, 17)]
nr_cols = [f'NR{i}' for i in range(1, 17)]
all_cols = pr_cols + nr_cols

# Define hyperparameter grid (reduced combinations to speed up execution)
param_grid = {
    'n_estimators': [100, 200],
    'max_depth': [None, 10],
    'min_samples_split': [2, 5],
    'max_features': ['sqrt', 'log2']
}

results = []
total_items = len(all_cols)

for idx, col in enumerate(all_cols, 1):
    if col not in data.columns:
        continue
    
    print(f"\nProcessing item {idx}/{total_items}: {col}")
    y = data[col]
    
    # Create KFold object
    kf = KFold(n_splits=5, shuffle=True, random_state=42)
    
    print("Performing grid search...")
    # Initialize GridSearchCV
    grid_search = GridSearchCV(
        estimator=RandomForestRegressor(random_state=42),
        param_grid=param_grid,
        cv=kf,
        scoring='neg_mean_squared_error',
        n_jobs=-1
    )
    
    # Fit model
    grid_search.fit(X, y)
    print("Grid search completed. Best parameters:", grid_search.best_params_)
    
    # Get best model
    best_rf = grid_search.best_estimator_
    
    # Calculate training set performance
    y_pred_train = best_rf.predict(X)
    r2_train = r2_score(y, y_pred_train)
    mse_train = mean_squared_error(y, y_pred_train)
    mae_train = mean_absolute_error(y, y_pred_train)
    
    print("Calculating cross-validation scores...")
    # Calculate cross-validation performance
    cv_scores = cross_val_score(best_rf, X, y, cv=kf, scoring='neg_mean_squared_error')
    mse_cv = -np.mean(cv_scores)
    
    # Get feature importance
    importances = best_rf.feature_importances_
    
    results.append([
        col, 
        grid_search.best_params_,
        r2_train, mse_train, mae_train, mse_cv,
        importances[0], importances[1], importances[2], importances[3]
    ])
    print(f"Completed processing {col}")

print("\nSaving results...")
# Create results DataFrame
results_df = pd.DataFrame(results, columns=[
    'DV', 'Best_Params',
    'R2_train', 'MSE_train', 'MAE_train', 'MSE_cv',
    'ScomFactor1_importance', 'ScomFactor2_importance',
    'ScapFactor1_importance', 'ScapFactor2_importance'
])

# Save results
os.makedirs('results', exist_ok=True)
results_df.to_csv('results/itemwise_rf_results.csv', index=False)

# Create Word document
try:
    doc = Document()
    
    # Add title
    doc.add_heading('Random Forest Regression Results', 0)
    
    # Add table
    table = doc.add_table(rows=1, cols=len(results_df.columns))
    table.style = 'Table Grid'
    
    # Add headers
    for i, col in enumerate(results_df.columns):
        cell = table.cell(0, i)
        cell.text = str(col)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add data
    for _, row in results_df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = f'{val:.4f}' if isinstance(val, float) else str(val)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Try to save document
    try:
        doc.save('results/itemwise_rf_results.docx')
    except PermissionError:
        import datetime
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_filename = f'results/itemwise_rf_results_{timestamp}.docx'
        print(f"\nWarning: Could not save to original file. Saving to {backup_filename} instead.")
        doc.save(backup_filename)
    
    print('\nRandom Forest regression results have been saved as CSV and Word files in the results folder.')
except Exception as e:
    print(f"\nError while saving Word document: {str(e)}")
    print("CSV file has been saved successfully. You can open it with Excel.") 